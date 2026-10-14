"""Gerador da planilha de rotulagem independente da amostra-ouro para o avaliador Kayque (LAB03S02).

Dois comandos:
  gerar    base.csv  -> planilha .xlsx com listas suspensas (só aceita valores válidos)
  exportar planilha.xlsx -> CSV no formato exato aceito por `validar_rotulagem`

A planilha tem duas abas: "Repositórios" (tipo_projeto e entrega_real, preenchidos uma
vez por repositório) e "Releases" (is_corretiva e justificativa, por release). Nenhum
rótulo é preenchido automaticamente: todo julgamento é do avaliador.

Uso (na raiz do repositório):
    python -m lab03.src.gerar_rotulagem_kayque gerar lab03/data/amostra_ouro_base.csv \
        --saida lab03/data/planilha_rotulagem_kayque.xlsx
    python -m lab03.src.gerar_rotulagem_kayque exportar lab03/data/planilha_rotulagem_kayque.xlsx \
        --saida lab03/data/rotulagem_amostra_ouro_kayque.csv --base lab03/data/amostra_ouro_base.csv
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import List, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .validar_rotulagem import (
    COLUNAS_BASE, COLUNAS_ROTULO, CORRETIVAS, ENTREGAS, TIPOS, validar,
)

ABA_REPOS = "Repositórios"
ABA_RELEASES = "Releases"
COLS_REPO = ["repo_id", "full_name", "primary_language", "repo_url", "tipo_projeto", "entrega_real", "observacao_repo"]
PREENCHER = PatternFill("solid", fgColor="FFF2CC")
CABECALHO = PatternFill("solid", fgColor="D9E1F2")


def _lista(valores) -> str:
    return '"' + ",".join(sorted(valores)) + '"'


def _estilo_cabecalho(ws, larguras: List[int]) -> None:
    for c, larg in enumerate(larguras, start=1):
        cel = ws.cell(row=1, column=c)
        cel.font = Font(bold=True)
        cel.fill = CABECALHO
        ws.column_dimensions[cel.column_letter].width = larg
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def _validacao(ws, coluna: int, valores, n_linhas: int) -> None:
    letra = ws.cell(row=1, column=coluna).column_letter
    dv = DataValidation(type="list", formula1=_lista(valores), allow_blank=True,
                        showErrorMessage=True, errorTitle="Valor inválido",
                        error="Escolha um valor da lista.")
    ws.add_data_validation(dv)
    dv.add(f"{letra}2:{letra}{n_linhas + 1}")
    for r in range(2, n_linhas + 2):
        ws.cell(row=r, column=coluna).fill = PREENCHER


def gerar(base_csv: Path, saida: Path) -> None:
    base = pd.read_csv(base_csv, encoding="utf-8").fillna("")
    wb = Workbook()

    ws = wb.active
    ws.title = ABA_REPOS
    ws.append(COLS_REPO)
    repos = base.drop_duplicates("repo_id").sort_values("repo_id")
    for _, r in repos.iterrows():
        ws.append([int(r["repo_id"]), r["full_name"], r["primary_language"], r["repo_url"], None, None, None])
    _estilo_cabecalho(ws, [8, 38, 16, 40, 24, 16, 50])
    _validacao(ws, 5, TIPOS, len(repos))
    _validacao(ws, 6, ENTREGAS, len(repos))

    wr = wb.create_sheet(ABA_RELEASES)
    colunas_info = [c for c in ["repo_id", "full_name", "release_num", "tag_name", "release_name",
                                "published_at", "release_url", "release_body"] if c in base.columns]
    wr.append(colunas_info + ["is_corretiva", "justificativa"])
    for _, r in base.sort_values(["repo_id", "release_num"]).iterrows():
        linha = [int(r[c]) if c in ("repo_id", "release_num") else r[c] for c in colunas_info]
        wr.append(linha + [None, None])
    larg = {"repo_id": 8, "full_name": 34, "release_num": 8, "tag_name": 16, "release_name": 36,
            "published_at": 22, "release_url": 45, "release_body": 60}
    _estilo_cabecalho(wr, [larg[c] for c in colunas_info] + [14, 60])
    n = len(base)
    _validacao(wr, len(colunas_info) + 1, CORRETIVAS, n)
    for r in range(2, n + 2):
        wr.cell(row=r, column=len(colunas_info) + 2).fill = PREENCHER
        url = wr.cell(row=r, column=colunas_info.index("release_url") + 1)
        if url.value:
            url.hyperlink = str(url.value)
            url.font = Font(color="0563C1", underline="single")
        if "release_body" in colunas_info:
            wr.cell(row=r, column=colunas_info.index("release_body") + 1).alignment = Alignment(wrap_text=True, vertical="top")

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    print(f"Planilha gerada: {saida} ({len(repos)} repositórios, {n} releases)")


def _ler_aba(wb, nome: str) -> pd.DataFrame:
    linhas = list(wb[nome].iter_rows(values_only=True))
    df = pd.DataFrame(linhas[1:], columns=list(linhas[0]))
    return df.dropna(how="all")


def exportar(planilha: Path, saida: Path, base_csv: Optional[Path]) -> List[str]:
    wb = load_workbook(planilha, data_only=True)
    repos = _ler_aba(wb, ABA_REPOS)
    rel = _ler_aba(wb, ABA_RELEASES)

    repos["repo_id"] = repos["repo_id"].astype(int)
    rel["repo_id"] = rel["repo_id"].astype(int)
    rel["release_num"] = rel["release_num"].astype(int)

    df = rel.merge(
        repos[["repo_id", "primary_language", "repo_url", "tipo_projeto", "entrega_real", "observacao_repo"]],
        on="repo_id", how="left",
    )
    df["justificativa"] = df["justificativa"].fillna("")
    sem_just = df["justificativa"].astype(str).str.strip() == ""
    df.loc[sem_just, "justificativa"] = df.loc[sem_just, "observacao_repo"].fillna("")
    df = df.fillna("")
    ordem = COLUNAS_BASE + (["release_body"] if "release_body" in df.columns else []) + COLUNAS_ROTULO
    df = df[ordem].sort_values(["repo_id", "release_num"])

    saida.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(saida, index=False, encoding="utf-8")
    print(f"CSV exportado: {saida} ({len(df)} linhas)")

    base = pd.read_csv(base_csv, encoding="utf-8") if base_csv and Path(base_csv).exists() else None
    return validar(df, base)


def main() -> int:
    p = argparse.ArgumentParser(description="Planilha de rotulagem da amostra-ouro")
    sub = p.add_subparsers(dest="cmd", required=True)
    g = sub.add_parser("gerar")
    g.add_argument("base")
    g.add_argument("--saida", required=True)
    e = sub.add_parser("exportar")
    e.add_argument("planilha")
    e.add_argument("--saida", required=True)
    e.add_argument("--base", default=None)
    a = p.parse_args()

    if a.cmd == "gerar":
        gerar(Path(a.base), Path(a.saida))
        return 0
    erros = exportar(Path(a.planilha), Path(a.saida), Path(a.base) if a.base else None)
    if erros:
        print(f"Atenção: o CSV ainda tem {len(erros)} pendência(s):")
        for x in erros:
            print(f" - {x}")
        return 1
    print("OK: rotulagem completa e válida.")
    return 0


if __name__ == "__main__":
    sys.exit(main())